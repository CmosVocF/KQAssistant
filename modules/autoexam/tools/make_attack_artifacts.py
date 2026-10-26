# -*- coding: utf-8 -*-
"""Create test artifacts used by the security probe (local only, throwaway)."""
import os
import openpyxl

ROOT = r'D:\Autoexam\tools\_deploytest'
VICTIM = os.path.join(ROOT, 'victim_docs')
os.makedirs(VICTIM, exist_ok=True)

# ---- 假“敏感”表格：一旦被 reload 就能通过 api/query 把内容读出来 ----
wb = openpyxl.Workbook()
ws = wb.active
ws.title = '选择题'
ws.append(['文件夹编号', '资源编号', '标题', '题目', '可否更改答案顺序', '答案',
           '选项 1', '选项 2', '选项 3', '选项 4', '分数', '难度', '状态', '解析'])
ws.append(['', '', '公司机密', '服务器管理员密码？', 'Y', 2,
           'password_abc123', 'letmein', 'HR@2024', 'qwerty', 1, 1, 'ON', '内部文件'])
ws.append(['', '', '公司机密', 'CEO 银行账号？', 'Y', 1,
           '6222 8888 6666 0001', '6222000000', '6217 0000', '6214 9999', 1, 1, 'ON', '内部文件'])
wb.save(os.path.join(VICTIM, '机密人事表.xlsx'))

# 第二份“隐藏”敏感数据，验证能枚举路径下多个文件
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = '判断题'
ws2.append(['文件夹编号', '资源编号', '标题', '题目', '答案', '分数', '难度', '状态', '解析'])
ws2.append(['', '', '服务器', 'SSH 密钥存放在 /root/.ssh/id_rsa？', 'True', 1, 1, 'ON', '内部'])
wb2.save(os.path.join(VICTIM, 'inventory.xlsx'))

# ---- 用一个大 xlsx 做 DoS 靶子（真实数据 120k 行，openpyxl 读起来很慢）----
wb3 = openpyxl.Workbook(write_only=True)
ws3 = wb3.create_sheet('选择题')
ws3.append(['文件夹编号', '资源编号', '标题', '题目', '可否更改答案顺序', '答案',
            '选项 1', '选项 2', '分数', '难度', '状态', '解析'])
for i in range(120000):
    ws3.append(['', '', 'DoS靶子', f'这是第 {i} 行用于拖慢加载的题目？（  ）。', 'Y', 1,
                f'选项A{i}', f'选项B{i}', 1, 1, 'ON', ''])
wb3.save(os.path.join(ROOT, 'dos_big.xlsx'))

print('artifacts:')
for f in os.listdir(ROOT):
    p = os.path.join(ROOT, f)
    print(' ', f, os.path.getsize(p) // 1024, 'KB')
